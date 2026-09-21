from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.adapters.reports.excel import ExcelReportExporter
from app.adapters.reports.memory_store import InMemoryReportStore
from app.application.reconciliation import ReconciliationService
from app.domain.entities import InvoiceBatch, PaymentBatch
from app.domain.matching.engine import MatchingEngine
from app.domain.matching.rules import ExplicitInvoiceRule
from app.domain.validator import ReconciliationValidator
from app.ports.reports import ReportTooLargeError
from tests.conftest import make_invoice, make_payment


def sample_result():  # type: ignore[no-untyped-def]
    return ReconciliationService(
        MatchingEngine((ExplicitInvoiceRule(),)), ReconciliationValidator()
    ).execute(
        PaymentBatch((make_payment(amount="110"),), ()),
        InvoiceBatch((make_invoice(),)),
    )


def test_excel_export_contains_required_sheets_and_values() -> None:
    report = ExcelReportExporter().export(sample_result())
    workbook = load_workbook(filename=__import__("io").BytesIO(report), data_only=True)
    assert workbook.sheetnames == ["Summary", "Invoices", "Unallocated", "Payments"]
    assert workbook["Summary"]["B10"].value == Decimal("10.00")
    assert workbook["Invoices"]["A2"].value == "100"
    assert workbook["Unallocated"]["E2"].value == Decimal("10.00")
    assert workbook["Payments"].freeze_panes == "A2"


def test_report_store_returns_report_and_enforces_capacity() -> None:
    store = InMemoryReportStore(max_reports=1)
    old = store.put(b"old")
    new = store.put(b"new")
    assert store.get(old) is None
    assert store.get(new) == b"new"


def test_report_store_expires_report() -> None:
    store = InMemoryReportStore(ttl_seconds=0)
    report_id = store.put(b"report")
    assert store.get(report_id) is None


def test_report_store_enforces_total_byte_capacity() -> None:
    store = InMemoryReportStore(max_reports=10, max_total_bytes=5)
    old = store.put(b"old")
    new = store.put(b"new")
    assert store.get(old) is None
    assert store.get(new) == b"new"


def test_report_store_rejects_single_report_over_byte_capacity() -> None:
    store = InMemoryReportStore(max_total_bytes=5)
    with pytest.raises(ReportTooLargeError, match="превышает допустимый размер"):
        store.put(b"report")
