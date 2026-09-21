from collections.abc import Iterable
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.adapters.parsing import normalize_identifier, parse_date, parse_money
from app.application.errors import ReconciliationError
from app.domain.entities import Invoice, InvoiceBatch


class ExcelInvoiceSource:
    REQUIRED_HEADERS = ("№ счета", "Дата счета", "Покупатель", "ИНН", "Сумма")

    def load(self, source: BinaryIO) -> InvoiceBatch:
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise ReconciliationError("Не удалось прочитать реестр счетов XLSX") from exc
        try:
            for worksheet in workbook.worksheets:
                header_row = self._find_header(worksheet.iter_rows(values_only=True))
                if header_row is not None:
                    return self._load_sheet(worksheet, header_row)
        finally:
            workbook.close()
        raise ReconciliationError("В реестре не найдена таблица счетов")

    def _find_header(self, rows: Iterable[tuple[object, ...]]) -> int | None:
        for index, row in enumerate(rows, start=1):
            values = tuple(str(value).strip() if value is not None else "" for value in row[:5])
            if values == self.REQUIRED_HEADERS:
                return index
        return None

    def _load_sheet(self, worksheet: Worksheet, header_row: int) -> InvoiceBatch:
        invoices: list[Invoice] = []
        seen: set[str] = set()
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            if str(row[0]).strip().casefold() == "итого":
                break
            number = normalize_identifier(row[0], f"№ счета, строка {row_number}")
            if number in seen:
                raise ReconciliationError(f"Номер счёта {number} встречается более одного раза")
            seen.add(number)
            invoices.append(
                Invoice(
                    number=number,
                    date=parse_date(row[1], f"Дата счёта, строка {row_number}"),
                    customer_name=normalize_identifier(
                        row[2], f"Покупатель, строка {row_number}"
                    ),
                    customer_inn=normalize_identifier(row[3], f"ИНН, строка {row_number}"),
                    amount=parse_money(row[4], f"Сумма, строка {row_number}"),
                )
            )
        if not invoices:
            raise ReconciliationError("В реестре не найдено ни одного счёта")
        return InvoiceBatch(tuple(invoices))
