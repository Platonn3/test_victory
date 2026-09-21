import io
from collections.abc import Iterable
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.entities import ReconciliationResult

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=16, bold=True, color="1F4E78")
_MONEY_FORMAT = '#,##0.00" ₽"'
_DATE_FORMAT = "dd.mm.yyyy"


class ExcelReportExporter:
    def export(self, result: ReconciliationResult) -> bytes:
        workbook = Workbook()
        summary = workbook.active
        assert isinstance(summary, Worksheet)
        summary.title = "Summary"
        self._summary(summary, result)
        self._invoices(workbook.create_sheet("Invoices"), result)
        self._unallocated(workbook.create_sheet("Unallocated"), result)
        self._payments(workbook.create_sheet("Payments"), result)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _summary(self, sheet: Worksheet, result: ReconciliationResult) -> None:
        sheet.append(["Сверка банковских поступлений"])
        sheet["A1"].font = _TITLE_FONT
        sheet.append([])
        rows: list[tuple[str, object]] = [
            ("Счетов", len(result.invoices)),
            ("Полностью оплачено", result.fully_paid_count),
            ("Частично оплачено", result.partially_paid_count),
            ("Не оплачено", result.unpaid_count),
            ("Сумма счетов", result.total_invoice_amount),
            ("Поступления", result.total_income),
            ("Распределено", result.total_allocated),
            ("Не распределено", result.total_unallocated),
            ("Остаток по счетам", result.total_outstanding),
            ("Проверка банковского баланса", result.balance_check),
        ]
        for label, value in rows:
            sheet.append([label, value])
        for row in range(7, 12):
            sheet.cell(row=row, column=2).number_format = _MONEY_FORMAT
        if result.warnings:
            sheet.append([])
            sheet.append(["Предупреждения"])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
            for warning in result.warnings:
                sheet.append([warning])
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 24

    def _invoices(self, sheet: Worksheet, result: ReconciliationResult) -> None:
        headers = (
            "invoice_number",
            "date",
            "customer",
            "inn",
            "invoice_amount",
            "allocated_amount",
            "outstanding_amount",
            "status",
        )
        rows = (
            (
                item.invoice.number,
                item.invoice.date,
                item.invoice.customer_name,
                item.invoice.customer_inn,
                item.invoice.amount,
                item.allocated_amount,
                item.outstanding_amount,
                item.status.value,
            )
            for item in result.invoices
        )
        self._table(sheet, headers, rows, money_columns=(5, 6, 7), date_columns=(2,))

    def _unallocated(self, sheet: Worksheet, result: ReconciliationResult) -> None:
        headers = (
            "date",
            "document_number",
            "payer",
            "inn",
            "amount",
            "purpose",
            "reason",
            "candidate_invoices",
        )
        rows = (
            (
                item.payment.date,
                item.payment.document_number,
                item.payment.payer_name,
                item.payment.payer_inn or "",
                item.amount,
                item.payment.purpose,
                item.reason.value,
                ", ".join(item.candidate_invoices),
            )
            for item in result.unallocated
        )
        self._table(sheet, headers, rows, money_columns=(5,), date_columns=(1,))

    def _payments(self, sheet: Worksheet, result: ReconciliationResult) -> None:
        headers = (
            "date",
            "document_number",
            "payer",
            "amount",
            "allocated_amount",
            "unallocated_amount",
            "invoice_numbers",
            "reason",
            "warnings",
        )
        rows: list[tuple[str | int | Decimal | date | None, ...]] = []
        for item in result.payments:
            rows.append(
                (
                    item.payment.date,
                    item.payment.document_number,
                    item.payment.payer_name,
                    item.payment.amount,
                    sum(allocation.amount for allocation in item.allocations),
                    item.unallocated_amount,
                    ", ".join(allocation.invoice_number for allocation in item.allocations),
                    item.reason.value,
                    "; ".join(item.warnings),
                )
            )
        for duplicate in result.duplicates:
            rows.append(
                (
                    "",
                    duplicate.document_number,
                    "",
                    duplicate.amount,
                    Decimal("0.00"),
                    Decimal("0.00"),
                    "",
                    "duplicate",
                    "Строка исключена из итогов",
                )
            )
        self._table(sheet, headers, rows, money_columns=(4, 5, 6), date_columns=(1,))

    @staticmethod
    def _table(
        sheet: Worksheet,
        headers: tuple[str, ...],
        rows: Iterable[tuple[str | int | Decimal | date | None, ...]],
        money_columns: tuple[int, ...],
        date_columns: tuple[int, ...],
    ) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        for data_row in rows:
            sheet.append(data_row)
        for row_index in range(2, sheet.max_row + 1):
            for column in money_columns:
                sheet.cell(row=row_index, column=column).number_format = _MONEY_FORMAT
            for column in date_columns:
                if isinstance(sheet.cell(row=row_index, column=column).value, date):
                    sheet.cell(row=row_index, column=column).number_format = _DATE_FORMAT
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(
            fitToPage=True, autoPageBreaks=False
        )
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.page_margins = PageMargins(
            left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2
        )
        for index, header in enumerate(headers, start=1):
            content_width = max(
                (
                    len(str(sheet.cell(row=row_index, column=index).value or ""))
                    for row_index in range(1, sheet.max_row + 1)
                ),
                default=len(header),
            )
            width = min(max(content_width + 2, len(header) + 2, 12), 45)
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
